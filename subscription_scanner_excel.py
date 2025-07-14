#!/usr/bin/env python3
"""
Export subscription scan results to Excel with comprehensive analysis
"""

import sqlite3
import re
from datetime import datetime, timedelta
from pathlib import Path
from collections import defaultdict
from difflib import SequenceMatcher
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

# Use the same keywords and logic from advanced scanner
# Import shared components from advanced scanner
from pathlib import Path
import sys
sys.path.insert(0, str(Path(__file__).parent))
from advanced_subscription_scanner import AdvancedSubscriptionScanner

class ExcelExporter(AdvancedSubscriptionScanner):
    """Export subscription data to Excel with multiple analysis sheets"""
    
    def export_to_excel(self, filename: str = None):
        """Export comprehensive analysis to Excel"""
        if not self.subscriptions:
            print("No data to export. Run scan_emails() first.")
            return
            
        if not filename:
            filename = f"subscriptions_analysis_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        
        filepath = Path.home() / "Desktop" / filename
        
        # Create workbook
        wb = Workbook()
        wb.remove(wb.active)  # Remove default sheet
        
        # 1. Summary Dashboard
        self._create_summary_sheet(wb)
        
        # 2. All Subscriptions
        self._create_subscriptions_sheet(wb)
        
        # 3. Duplicate Analysis
        self._create_duplicates_sheet(wb)
        
        # 4. Account Breakdown
        self._create_account_sheet(wb)
        
        # 5. Payment Issues
        self._create_issues_sheet(wb)
        
        # 6. Timeline Analysis
        self._create_timeline_sheet(wb)
        
        # Save workbook
        wb.save(filepath)
        print(f"\n✅ Excel report saved to: {filepath}")
        return filepath
    
    def _create_summary_sheet(self, wb):
        """Create summary dashboard"""
        ws = wb.create_sheet("Summary Dashboard")
        
        # Title
        ws['A1'] = "Subscription Analysis Summary"
        ws['A1'].font = Font(size=16, bold=True)
        ws['A2'] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
        
        # Key metrics
        row = 4
        metrics = [
            ("Total Unique Subscriptions", len(self.subscriptions)),
            ("Active Subscriptions", sum(1 for d in self.subscriptions.values() if hasattr(d, 'status') and d.status == 'active')),
            ("Cancelled Subscriptions", sum(1 for d in self.subscriptions.values() if hasattr(d, 'status') and d.status == 'cancelled')),
            ("Trial Subscriptions", sum(1 for d in self.subscriptions.values() if hasattr(d, 'status') and d.status == 'trial')),
            ("Payment Issues", sum(1 for d in self.subscriptions.values() if hasattr(d, 'status') and d.status in ['failed', 'past_due'])),
        ]
        
        ws['A4'] = "KEY METRICS"
        ws['A4'].font = Font(bold=True)
        
        for metric, value in metrics:
            row += 1
            ws[f'A{row}'] = metric
            ws[f'B{row}'] = value
        
        # Financial summary
        row += 2
        ws[f'A{row}'] = "FINANCIAL SUMMARY"
        ws[f'A{row}'].font = Font(bold=True)
        
        total_monthly = 0
        total_cancelled_savings = 0
        
        for company, data in self.subscriptions.items():
            # Handle both Subscription objects and dictionary format
            if hasattr(data, 'avg_amount'):
                avg_amount = data.avg_amount if data.avg_amount else 0
                status = data.status
            else:
                avg_amount = sum(data.get('amounts', [])) / len(data.get('amounts', [])) if data.get('amounts') else 0
                status = data.get('status', 'unknown')
            
            if status == 'active' and avg_amount > 0:
                total_monthly += avg_amount
            elif status == 'cancelled' and avg_amount > 0:
                total_cancelled_savings += avg_amount
        
        financial_metrics = [
            ("Estimated Monthly Active", f"${total_monthly:.2f}"),
            ("Estimated Annual Active", f"${total_monthly * 12:.2f}"),
            ("Monthly Cancelled Savings", f"${total_cancelled_savings:.2f}"),
            ("Annual Cancelled Savings", f"${total_cancelled_savings * 12:.2f}"),
        ]
        
        for metric, value in financial_metrics:
            row += 1
            ws[f'A{row}'] = metric
            ws[f'B{row}'] = value
        
        # Duplicate waste
        duplicate_groups = self._find_duplicates()
        duplicate_waste = 0
        
        for group in duplicate_groups:
            group_accounts = set()
            group_total = 0
            
            for company in group:
                data = self.subscriptions[company]
                group_accounts.update(data.accounts)
                if data.amounts:
                    group_total += data.avg_amount if data.avg_amount else 0
            
            if len(group_accounts) > 1:
                duplicate_waste += group_total - (group_total / len(group))
        
        row += 2
        ws[f'A{row}'] = "DUPLICATE WASTE"
        ws[f'A{row}'].font = Font(bold=True)
        row += 1
        ws[f'A{row}'] = "Monthly Duplicate Waste"
        ws[f'B{row}'] = f"${duplicate_waste:.2f}"
        row += 1
        ws[f'A{row}'] = "Annual Duplicate Waste"
        ws[f'B{row}'] = f"${duplicate_waste * 12:.2f}"
        
        # Format columns
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 15
    
    def _create_subscriptions_sheet(self, wb):
        """Create detailed subscriptions list"""
        ws = wb.create_sheet("All Subscriptions")
        
        # Prepare data
        rows = []
        for company, data in self.subscriptions.items():
            avg_amount = data.avg_amount if data.avg_amount else 0
            last_date = data.last_seen
            first_date = data.first_seen
            
            rows.append({
                'Company': company,
                'Status': data.status.upper(),
                'Accounts': ', '.join(data.accounts),
                'Account Count': len(data.accounts),
                'Avg Amount': avg_amount,
                'Annual Cost': avg_amount * 12,
                'Email Count': data.email_count,
                'First Seen': first_date.strftime('%Y-%m-%d') if first_date else '',
                'Last Seen': last_date.strftime('%Y-%m-%d') if last_date else '',
                'Days Since Last': (datetime.now() - last_date).days if last_date else '',
                'Payment Methods': 'N/A'  # Not available in current data structure
            })
        
        # Create DataFrame and add to sheet
        df = pd.DataFrame(rows)
        df = df.sort_values('Avg Amount', ascending=False)
        
        # Headers
        headers = list(df.columns)
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            cell.font = Font(color="FFFFFF", bold=True)
        
        # Data
        for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=False), 2):
            for c_idx, value in enumerate(row, 1):
                cell = ws.cell(row=r_idx, column=c_idx, value=value)
                
                # Color code by status
                if c_idx == 2:  # Status column
                    if value == 'ACTIVE':
                        cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                    elif value == 'CANCELLED':
                        cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                    elif value in ['FAILED', 'PAST_DUE']:
                        cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
                
                # Format currency
                if c_idx in [5, 6]:  # Amount columns
                    cell.number_format = '"$"#,##0.00'
        
        # Auto-fit columns
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
    
    def _create_duplicates_sheet(self, wb):
        """Create duplicate analysis sheet"""
        ws = wb.create_sheet("Duplicate Analysis")
        
        duplicate_groups = self._find_duplicates()
        
        if not duplicate_groups:
            ws['A1'] = "No duplicate subscriptions found!"
            return
        
        ws['A1'] = "Duplicate Subscription Analysis"
        ws['A1'].font = Font(size=14, bold=True)
        
        row = 3
        total_waste = 0
        
        for group_idx, group in enumerate(duplicate_groups, 1):
            ws[f'A{row}'] = f"Duplicate Group {group_idx}"
            ws[f'A{row}'].font = Font(bold=True)
            row += 1
            
            # Headers
            headers = ['Service', 'Accounts', 'Amount', 'Status', 'Last Seen']
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col, value=header)
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
            row += 1
            
            # Group data
            group_accounts = set()
            group_total = 0
            
            for company in group:
                data = self.subscriptions[company]
                group_accounts.update(data.accounts)
                
                avg_amount = data.avg_amount if data.avg_amount else 0
                last_date = data.last_seen
                
                ws[f'A{row}'] = company
                ws[f'B{row}'] = ', '.join(data.accounts)
                ws[f'C{row}'] = avg_amount
                ws[f'C{row}'].number_format = '"$"#,##0.00'
                ws[f'D{row}'] = data.status.upper()
                ws[f'E{row}'] = last_date.strftime('%Y-%m-%d') if last_date else ''
                
                group_total += avg_amount
                row += 1
            
            # Calculate waste
            if len(group_accounts) > 1:
                waste = group_total - (group_total / len(group))
                ws[f'A{row}'] = "Potential Monthly Waste:"
                ws[f'A{row}'].font = Font(bold=True)
                ws[f'C{row}'] = waste
                ws[f'C{row}'].number_format = '"$"#,##0.00'
                ws[f'C{row}'].font = Font(bold=True, color="FF0000")
                total_waste += waste
            
            row += 2
        
        # Total waste summary
        ws[f'A{row}'] = "TOTAL MONTHLY DUPLICATE WASTE:"
        ws[f'A{row}'].font = Font(size=12, bold=True)
        ws[f'C{row}'] = total_waste
        ws[f'C{row}'].number_format = '"$"#,##0.00'
        ws[f'C{row}'].font = Font(size=12, bold=True, color="FF0000")
        
        row += 1
        ws[f'A{row}'] = "TOTAL ANNUAL DUPLICATE WASTE:"
        ws[f'A{row}'].font = Font(size=12, bold=True)
        ws[f'C{row}'] = total_waste * 12
        ws[f'C{row}'].number_format = '"$"#,##0.00'
        ws[f'C{row}'].font = Font(size=12, bold=True, color="FF0000")
    
    def _create_account_sheet(self, wb):
        """Create account breakdown sheet"""
        ws = wb.create_sheet("By Account")
        
        # Aggregate by account
        account_data = defaultdict(lambda: {
            'subscriptions': [],
            'total': 0,
            'active_count': 0,
            'cancelled_count': 0
        })
        
        for company, data in self.subscriptions.items():
            avg_amount = sum(data['amounts']) / len(data['amounts']) if data['amounts'] else 0
            
            for account in data['accounts']:
                acc = account_data[account]
                acc['subscriptions'].append({
                    'company': company,
                    'amount': avg_amount,
                    'status': data['status']
                })
                
                if data['status'] == 'active':
                    acc['total'] += avg_amount
                    acc['active_count'] += 1
                elif data['status'] == 'cancelled':
                    acc['cancelled_count'] += 1
        
        # Create summary
        row = 1
        ws['A1'] = "Subscription Breakdown by Email Account"
        ws['A1'].font = Font(size=14, bold=True)
        
        row = 3
        
        for account, acc_data in sorted(account_data.items(), 
                                      key=lambda x: x[1]['total'], 
                                      reverse=True):
            # Account header
            ws[f'A{row}'] = account
            ws[f'A{row}'].font = Font(bold=True, size=12)
            ws[f'B{row}'] = f"Total Active: ${acc_data['total']:.2f}/mo"
            ws[f'C{row}'] = f"({acc_data['active_count']} active, {acc_data['cancelled_count']} cancelled)"
            row += 1
            
            # Subscription details
            for sub in sorted(acc_data['subscriptions'], 
                            key=lambda x: x['amount'], 
                            reverse=True):
                ws[f'B{row}'] = sub['company']
                ws[f'C{row}'] = sub['amount']
                ws[f'C{row}'].number_format = '"$"#,##0.00'
                ws[f'D{row}'] = sub['status'].upper()
                
                # Color code
                if sub['status'] == 'active':
                    ws[f'D{row}'].fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                elif sub['status'] == 'cancelled':
                    ws[f'D{row}'].fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                
                row += 1
            
            row += 1  # Space between accounts
    
    def _create_issues_sheet(self, wb):
        """Create payment issues sheet"""
        ws = wb.create_sheet("Payment Issues")
        
        issues = []
        
        for company, data in self.subscriptions.items():
            if data['status'] in ['failed', 'past_due']:
                avg_amount = sum(data['amounts']) / len(data['amounts']) if data['amounts'] else 0
                last_date = datetime.fromtimestamp(max(data['dates'])) if data['dates'] else None
                
                # Get most recent problem email
                problem_email = None
                if data['emails']:
                    for email in sorted(data['emails'], key=lambda x: x['date'], reverse=True):
                        if any(word in email['subject'].lower() for word in ['fail', 'declin', 'past due', 'overdue']):
                            problem_email = email
                            break
                
                issues.append({
                    'Company': company,
                    'Status': data['status'].upper(),
                    'Accounts': ', '.join(data['accounts']),
                    'Amount': avg_amount,
                    'Last Contact': last_date.strftime('%Y-%m-%d') if last_date else '',
                    'Days Ago': (datetime.now() - last_date).days if last_date else '',
                    'Problem Subject': problem_email['subject'][:60] if problem_email else '',
                    'Payment Method': ', '.join(data['payment_methods'])
                })
        
        if not issues:
            ws['A1'] = "No payment issues found!"
            ws['A1'].font = Font(size=12, color="008000")
            return
        
        ws['A1'] = "Subscriptions with Payment Issues"
        ws['A1'].font = Font(size=14, bold=True, color="FF0000")
        
        # Create DataFrame
        df = pd.DataFrame(issues)
        df = df.sort_values('Amount', ascending=False)
        
        # Add to sheet
        for col, header in enumerate(df.columns, 1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
            cell.font = Font(color="FFFFFF", bold=True)
        
        for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=False), 4):
            for c_idx, value in enumerate(row, 1):
                cell = ws.cell(row=r_idx, column=c_idx, value=value)
                if c_idx == 4:  # Amount column
                    cell.number_format = '"$"#,##0.00'
    
    def _create_timeline_sheet(self, wb):
        """Create timeline analysis sheet"""
        ws = wb.create_sheet("Timeline Analysis")
        
        ws['A1'] = "Subscription Activity Timeline"
        ws['A1'].font = Font(size=14, bold=True)
        
        # Categorize by last activity
        categories = {
            'Active (< 30 days)': [],
            'Recent (30-90 days)': [],
            'Inactive (90-180 days)': [],
            'Dormant (180-365 days)': [],
            'Ancient (> 365 days)': []
        }
        
        now = datetime.now()
        
        for company, data in self.subscriptions.items():
            if not data['dates']:
                continue
                
            last_date = datetime.fromtimestamp(max(data['dates']))
            days_ago = (now - last_date).days
            avg_amount = sum(data['amounts']) / len(data['amounts']) if data['amounts'] else 0
            
            item = {
                'company': company,
                'last_date': last_date,
                'days_ago': days_ago,
                'amount': avg_amount,
                'status': data['status'],
                'accounts': len(data['accounts'])
            }
            
            if days_ago < 30:
                categories['Active (< 30 days)'].append(item)
            elif days_ago < 90:
                categories['Recent (30-90 days)'].append(item)
            elif days_ago < 180:
                categories['Inactive (90-180 days)'].append(item)
            elif days_ago < 365:
                categories['Dormant (180-365 days)'].append(item)
            else:
                categories['Ancient (> 365 days)'].append(item)
        
        row = 3
        
        for category, items in categories.items():
            if not items:
                continue
                
            # Category header
            ws[f'A{row}'] = category
            ws[f'A{row}'].font = Font(bold=True, size=12)
            ws[f'B{row}'] = f"({len(items)} subscriptions)"
            row += 1
            
            # Headers
            headers = ['Company', 'Last Activity', 'Days Ago', 'Amount/mo', 'Status', 'Accounts']
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col, value=header)
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
            row += 1
            
            # Items
            for item in sorted(items, key=lambda x: x['days_ago']):
                ws[f'A{row}'] = item['company']
                ws[f'B{row}'] = item['last_date'].strftime('%Y-%m-%d')
                ws[f'C{row}'] = item['days_ago']
                ws[f'D{row}'] = item['amount']
                ws[f'D{row}'].number_format = '"$"#,##0.00'
                ws[f'E{row}'] = item['status'].upper()
                ws[f'F{row}'] = item['accounts']
                
                # Highlight potentially forgotten subscriptions
                if item['days_ago'] > 180 and item['status'] == 'active':
                    for col in range(1, 7):
                        ws.cell(row=row, column=col).fill = PatternFill(
                            start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"
                        )
                
                row += 1
            
            row += 1  # Space between categories

def main():
    print("🚀 Advanced Subscription Scanner with Excel Export")
    print("=" * 60)
    
    exporter = ExcelExporter()
    
    # Scan emails (20 years)
    exporter.scan_emails(days_back=7300)
    
    # Generate console report
    exporter.display_results()
    
    # Export to Excel
    print("\n📊 Generating Excel report...")
    filepath = exporter.export_to_excel()
    
    print(f"\n✨ Complete! Open the Excel file to see:")
    print("  • Summary dashboard with key metrics")
    print("  • All subscriptions with amounts and status")
    print("  • Duplicate analysis with waste calculations")
    print("  • Breakdown by email account")
    print("  • Payment issues that need attention")
    print("  • Timeline showing inactive subscriptions")

if __name__ == "__main__":
    main()
